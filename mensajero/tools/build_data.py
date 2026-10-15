#!/usr/bin/env python3
"""Planilla de Temas Macro (Excel) -> datos JSON inyectados en public/index.html.

Uso: python3 tools/build_data.py [ruta.xlsx]
Lee las hojas «Temas Macro AAAA», «Historia y Marcos», «Guia de Valores ADN» y, si existe, «Guiones AAAA».
El bloque entre /*DATA*/ y /*/DATA*/ de index.html se reemplaza; no editar a mano.
"""
import json, re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DIAS = ['Domingo', 'Lunes', 'Martes', 'Miercoles', 'Jueves', 'Viernes', 'Sabado']
DIAS_TXT = ['Domingo', 'Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado']


def txt(v):
    return '' if v is None else str(v).strip()


def puntos(v):
    # admite «• texto», «- texto» y «1. Título: explicación»
    return [re.sub(r'^(?:[•\-\s]+|\d+[.)]\s*)', '', l).strip() for l in txt(v).splitlines() if l.strip()]


def versiculo(v):
    # «Cita — texto» ; la cita queda separada para poder corregirla aparte
    s = txt(v)
    if '—' in s:
        ref, texto = s.split('—', 1)
        return {'ref': ref.strip(), 'texto': ' '.join(texto.split())}
    return {'ref': '', 'texto': ' '.join(s.split())}


def frase(v):
    s = txt(v)
    if '—' in s:
        f, autor = s.rsplit('—', 1)
        return {'texto': f.strip().strip('"“”'), 'autor': autor.strip()}
    return {'texto': s.strip('"“”'), 'autor': ''}


def quitar_tildes(s):
    import unicodedata
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = [txt(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if any(r)]


def revisar(anio, semanas):
    """Revisión de la planilla: avisa lo que conviene corregir antes de publicar.
    Escribe content/revision_planilla.txt; no detiene la generación."""
    avisos, usos = [], {}
    clave = lambda ref: re.sub(r'\s+', ' ', quitar_tildes(ref).lower().replace('—', '-')).strip()
    for w in semanas:
        for campo in ('enfoque', 'programa', 'locutor', 'historia', 'valor'):
            if not w[campo]:
                avisos.append(f'Semana {w["n"]}: falta «{campo}».')
        for i, st in enumerate(w['subtemas'], 1):
            if not st['titulo']:
                avisos.append(f'Semana {w["n"]}: falta el subtema {i}.')
        for i, d in enumerate(w['dias']):
            dia = f'Semana {w["n"]} · {DIAS_TXT[i]}'
            if not d['ref']:
                avisos.append(f'{dia}: el versículo no tiene el formato «Cita — texto».')
            else:
                usos.setdefault(clave(d['ref']), []).append(dia)
            if len(d['puntos']) != 3:
                avisos.append(f'{dia}: tiene {len(d["puntos"])} puntos (se esperan 3).')
    repetidos = [(ref, l) for ref, l in usos.items() if len(l) > 1]
    lineas = [f'Revisión de la planilla {anio}', '']
    lineas += ['VERSÍCULOS REPETIDOS'] + [f'- {ref}: ' + '; '.join(l) for ref, l in sorted(repetidos)] + ['']
    lineas += ['OTROS AVISOS'] + [f'- {a}' for a in avisos]
    (ROOT / 'content' / 'revision_planilla.txt').write_text('\n'.join(lineas) + '\n', encoding='utf-8')
    print(f'Revisión: {len(repetidos)} versículos repetidos · {len(avisos)} avisos → content/revision_planilla.txt')


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    hoja = next(n for n in wb.sheetnames if n.startswith('Temas Macro'))
    anio = int(re.search(r'\d{4}', hoja).group())

    marcos = {}
    for r in rows(wb['Historia y Marcos']):
        marcos[(txt(r['Tema']), int(r['Ronda']))] = {
            'historiaTitulo': txt(r['Historia Biblica - Titulo']),
            'historia': txt(r['Historia Biblica - Texto']),
            'teorico': txt(r['Marco Teorico - Afirmacion']),
            'fuente': txt(r['Marco Teorico - Fuente']),
            'claves': [s.strip() for s in txt(r['Marco Biblico - Versiculos clave']).split('|') if s.strip()],
            'contexto': txt(r['Marco Biblico - Contexto']),
            'conexion': txt(r['Marco Biblico - Conexion']),
        }

    valores = {}
    for r in rows(wb['Guia de Valores ADN']):
        valores[txt(r['Valor'])] = {
            'definicion': txt(r['Definicion aplicada']),
            'contenido': txt(r['Como se nota en el contenido']),
            'pregunta': txt(r['Pregunta guia']),
            'ejemplo': txt(r['Ejemplo de aplicacion']),
        }

    # guiones para el aire (opcional): una fila por día, lunes a sábado.
    # Lunes a viernes se apoyan en el versículo del día (sus 3 puntos son los del guion);
    # el sábado es el resumen de la semana.
    guiones = {}
    hg = f'Guiones {anio}'
    if hg in wb.sheetnames:
        for r in rows(wb[hg]):
            dia = next(i for i, d in enumerate(DIAS) if d.lower() == quitar_tildes(txt(r['Dia'])).lower())
            g = {'titulo': txt(r['Titulo']),
                 'dato': {'texto': txt(r['Dato curioso']), 'fuente': txt(r['Fuente dato'])}}
            if dia == 6:  # sábado = resumen: Desarrollo 1-3 son las ideas clave de cada subtema, Solucion 1 la acción
                g |= {'tipo': 'resumen', 'ideas': [txt(r[f'Desarrollo {i}']) for i in (1, 2, 3)], 'accion': txt(r['Solucion 1'])}
            else:
                g['puntos'] = [{'desarrollo': txt(r[f'Desarrollo {i}']), 'problema': txt(r[f'Problematica {i}']),
                                'solucion': txt(r[f'Solucion {i}'])} for i in (1, 2, 3)]
            guiones[(int(r['Semana']), dia)] = g

    semanas = []
    for r in rows(wb[hoja]):
        d, m, a = txt(r['Fecha inicio (Dom)']).split('/')
        tema, ronda = txt(r['Tema']), int(r['Ronda'])
        semanas.append({
            'n': int(r['Semana']),
            'inicio': f'{a}-{m}-{d}',
            'color': txt(r['Color']),
            'tema': tema,
            'ronda': ronda,
            'enfoque': txt(r['Enfoque']),
            'subtemas': [{
                'titulo': txt(r[f'Subtema {i}']),
                'desarrollo': txt(r[f'Desarrollo {i}']),
                'nota': txt(r[f'Nota Locutor {i}']),
                'frase': frase(r[f'Frase Celebre {i}']),
            } for i in (1, 2, 3)],
            'historia': txt(r['Historia Ilustrativa']),
            'valor': txt(r['Valor (ADN)']),
            'programa': txt(r['Programa']),
            'locutor': txt(r['Locutor']),
            'dias': [{**versiculo(r[f'Versiculo {d}']), 'puntos': puntos(r[f'3 Puntos {d}']),
                      'guion': guiones.get((int(r['Semana']), i))} for i, d in enumerate(DIAS)],
            'marco': marcos.get((tema, ronda)),
        })

    # información institucional (opcional): hoja «Mensajero» con Seccion, Titulo, Texto, Para el aire
    mensajero = []
    if 'Mensajero' in wb.sheetnames:
        for r in rows(wb['Mensajero']):
            texto = txt(r['Texto'])
            mensajero.append({'seccion': txt(r['Seccion']), 'titulo': txt(r['Titulo']),
                              'texto': '' if texto.startswith('(por completar') else texto,
                              'aire': txt(r.get('Para el aire')).lower().startswith('s')})

    data = {'anio': anio, 'semanas': semanas, 'valores': valores, 'mensajero': mensajero}
    revisar(anio, semanas)
    html = ROOT / 'public' / 'index.html'
    src = html.read_text(encoding='utf-8')
    blob = json.dumps(data, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
    out, n = re.subn(r'/\*DATA\*/.*?/\*/DATA\*/', lambda _: f'/*DATA*/{blob}/*/DATA*/', src, flags=re.S)
    if n != 1:
        sys.exit('No se encontró el marcador /*DATA*/ en index.html')
    html.write_text(out, encoding='utf-8')
    sin_marco = [s['n'] for s in semanas if not s['marco']]
    print(f'{anio}: {len(semanas)} semanas · {len(marcos)} marcos · {len(valores)} valores · {len(guiones)} días con guion'
          + (f' · sin marco: {sin_marco}' if sin_marco else ''))


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else next((ROOT / 'content' / 'planillas').glob('*.xlsx')))
