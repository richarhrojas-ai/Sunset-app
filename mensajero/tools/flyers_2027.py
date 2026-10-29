# Planilla para los diseñadores de flyers: un versículo por día (364), con fecha, tema, color de la semana y fechas especiales.
# Lee los datos ya cargados en la app (public/index.html), así la planilla siempre coincide con lo que ven los locutores.
# Uso: python3 tools/flyers_2027.py  →  content/Versiculos_flyers_2027.xlsx  (correr después de build_data.py)
import json, re, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
HTML = ROOT / 'public' / 'index.html'
SALIDA = ROOT / 'content' / 'Versiculos_flyers_2027.xlsx'
DATA = json.loads(re.search(r'/\*DATA\*/(.*?)/\*/DATA\*/', HTML.read_text(encoding='utf-8'), re.S).group(1))

LIBROS = {'Genesis': 'Génesis', 'Galatas': 'Gálatas', 'Exodo': 'Éxodo', 'Josue': 'Josué', 'Jeremias': 'Jeremías', 'Isaias': 'Isaías',
          'Eclesiastes': 'Eclesiastés', 'Levitico': 'Levítico', 'Nehemias': 'Nehemías', 'Salmo': 'Salmos', 'Numeros': 'Números',
          'Cronicas': 'Crónicas', 'Malaquias': 'Malaquías'}
ESPECIALES = {'2027-03-26': 'Viernes Santo', '2027-03-28': 'Domingo de Pascua', '2027-12-25': 'Navidad'}
DIAS = ['Domingo', 'Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado']
MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
TEMAS = {'Produccion y Desarrollo': 'Producción y Desarrollo'}
LARGO = 320  # caracteres: de ahí para arriba el flyer necesita letra más chica o formato vertical


def cita(ref):
    m = re.match(r'^(\d )?(\S+)(.*)$', ref)
    return (m.group(1) or '') + LIBROS.get(m.group(2), m.group(2)) + m.group(3)


def para_flyer(t):
    """Mismo texto RVC; solo se arreglan los bordes del fragmento: comillas sin pareja y cierre con coma o punto y coma."""
    t = re.sub(r'[ \t]+', ' ', t).strip()
    for abre, cierra in (('«', '»'), ('“', '”')):
        if t.startswith(cierra) and t.count(cierra) > t.count(abre):
            t = t[1:].lstrip()
        if t.count(cierra) > t.count(abre):
            k = t.rfind(cierra)
            t = (t[:k] + t[k + 1:]).rstrip()
        if t.count(abre) > t.count(cierra):
            t = re.sub(r'[,;:]$', '.', t.rstrip())
            if not re.search(r'[.!?]$', t):
                t += '.'
            t += cierra
    t = re.sub(r'\.\.$', '.', re.sub(r'[,;:]$', '.', t.rstrip()))
    return t if re.search(r'[.!?»”]$', t) else t + '.'


wb = openpyxl.Workbook()
ind = wb.active
ind.title = 'Indicaciones'
for linea in [
    'Versículos para flyers 2027 · Mensajero 100.9 FM · Edificando una Cultura', '',
    'Qué hay: un versículo por día, del domingo 3 de enero de 2027 al sábado 1 de enero de 2028 (364 flyers), en la hoja «Flyers 2027».',
    'Son los versículos definitivos: los mismos que usa la app de los locutores, con los cambios aprobados el 23 de octubre de 2026.', '',
    'Para el diseño:',
    '• Copiar el texto tal cual: es la Reina Valera Contemporánea (RVC). No cambiar palabras, no resumir y no pasar a otra versión.',
    '• La cita va siempre visible (libro, capítulo y versículos).',
    '• Créditos en cada flyer, en letra chica: «Reina Valera Contemporánea® © Sociedades Bíblicas Unidas, 2009, 2011. Usado con permiso de la Sociedad Bíblica del Paraguay.»',
    '• Los saltos de línea dentro de algunos versículos (salmos, proverbios) son del texto poético: se pueden respetar o unir en una sola línea.',
    '• Fechas especiales (filas amarillas): Viernes Santo (26 de marzo), Domingo de Pascua (28 de marzo) y Navidad (25 de diciembre). Merecen un diseño especial.',
    f'• Columna «Caracteres»: en naranja los textos de más de {LARGO} caracteres; necesitan letra más chica o formato vertical.',
    '• Color de la semana: es el color de la serie de esa semana en la app; sirve como guía de paleta.', '',
    'Ante cualquier duda sobre un texto, consultar antes de publicar: un error en un versículo se nota más que cualquier detalle de diseño.']:
    ind.append([linea])
ind['A1'].font = Font(bold=True, size=14, color='1F3A5F')
ind.column_dimensions['A'].width = 140
for fila in ind.iter_rows():
    fila[0].alignment = Alignment(wrap_text=True, vertical='top')

ws = wb.create_sheet('Flyers 2027')
ws.append(['N.º', 'Fecha', 'Día', 'Semana', 'Tema de la semana', 'Color de la semana', 'Fecha especial', 'Cita', 'Texto del versículo (RVC)', 'Caracteres'])
n = 0
for w in DATA['semanas']:
    inicio = datetime.date.fromisoformat(w['inicio'])
    for i, d in enumerate(w['dias']):
        n += 1
        f = inicio + datetime.timedelta(i)
        texto = para_flyer(d['texto'])
        ws.append([n, f'{f.day} de {MESES[f.month - 1]} de {f.year}', DIAS[i], w['n'], TEMAS.get(w['tema'], w['tema']), w['color'],
                   ESPECIALES.get(f.isoformat(), ''), cita(d['ref']), texto, len(texto)])

linea = Side(style='thin', color='CCCCCC')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F3A5F')
    c.alignment = Alignment(vertical='center', wrap_text=True)
for fila in ws.iter_rows(min_row=2):
    for c in fila:
        c.alignment = Alignment(vertical='top', wrap_text=True)
        c.border = Border(bottom=linea)
        if fila[6].value:
            c.fill = PatternFill('solid', fgColor='FFF4CC')
    fila[6].font = Font(bold=True)
    fila[7].font = Font(bold=True)
    if fila[9].value > LARGO:
        fila[9].fill = PatternFill('solid', fgColor='FFE0CC')
for k, ancho in enumerate([6, 20, 11, 8, 24, 14, 16, 22, 90, 11], 1):
    ws.column_dimensions[get_column_letter(k)].width = ancho
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions
wb.active = 1
wb.save(SALIDA)
print(f'{n} versículos en {SALIDA.relative_to(ROOT)}')
