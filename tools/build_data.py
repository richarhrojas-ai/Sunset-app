#!/usr/bin/env python3
"""Convierte la planilla de contenido en data/sunset-2027.js.

Uso:  python3 tools/build_data.py [ruta.xlsx]

Lee las hojas "Frases" (52 semanas), "Preguntas" (49 palabras × 5 preguntas)
y, si existe, "Mapeo Puente" (referencia interna de El Puente Relacional),
valida que las 3 palabras principales de cada semana coincidan con la rotación
calculada y escribe un único archivo JS que la app carga con <script>.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "sunset-contenido-2027.xlsx"
OUT = ROOT / "data" / "sunset-2027.js"

YEAR = 2027
FIRST_SUNDAY = "2027-01-03"

DIMENSIONS = ["Regalo de Dios", "Prioridades", "Temas de Interés", "Desafío",
              "Valores", "Estrategia", "Sistema"]

FRASES_COLS = {
    "color": "Color",
    "week": "Semana",
    "dates": "Fecha",
    "p1": "Principal 1",
    "p2": "Principal 2",
    "p3": "Principal 3",
    "title": "Domingo — Título",
    "story": "Domingo (Historia)",
    "hook": "Domingo — Frase gancho",
    "regalo": "@Regalo (Lunes)",
    "pensamiento": "@Pensamiento (Martes)",
    "desafio": "@Desafío (Miércoles)",
    "consejo": "@Consejo (Jueves)",
    "frase": "@Frase (Viernes)",
    "sabado": "Sábado (Evaluación)",
    "status": "Estado",
    "refs": "Referencias internas",
    "design": "Guía de diseño (paleta + imagen)",
}


def norm(word):
    """Unifica variantes de escritura entre hojas ("V.Cristiana" / "V. Cristiana")."""
    return re.sub(r"[\s.]", "", word or "").lower()


def clean(value):
    return str(value).strip() if value is not None else ""


def split_sabado(text):
    parts = re.split(r"\s*\d\)\s*", text)
    return [p.strip() for p in parts if p.strip()]


def parse_palette(design):
    return re.findall(r"#[0-9A-Fa-f]{6}", design.split("\n")[0])[:3]


def parse_image(design):
    m = re.search(r"Imagen sugerida:\s*(.+)", design)
    return m.group(1).strip() if m else ""


def read_puente(wb):
    """Hoja opcional: semana → etapa, eje, factor y nota de uso."""
    if "Mapeo Puente" not in wb.sheetnames:
        return {}
    ws = wb["Mapeo Puente"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and clean(r[0]) == "Semana")
    header = [clean(c) for c in rows[hdr_i]]
    col = lambda name: next(i for i, h in enumerate(header) if h.startswith(name))
    c_stage, c_axis, c_factor, c_note = col("Etapa"), col("Eje"), col("Factor"), col("Nota")
    out = {}
    for r in rows[hdr_i + 1:]:
        if not isinstance(r[0], (int, float)):
            continue
        out[int(r[0])] = {
            "stage": clean(r[c_stage]),
            "axis": clean(r[c_axis]),
            "factor": clean(r[c_factor]),
            "note": clean(r[c_note]),
        }
    return out


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    puente = read_puente(wb)

    # ── Tablero 7×7 ──
    board = [[] for _ in DIMENSIONS]
    questions = {}
    ws = wb["Preguntas"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]:
            continue
        dim, word = clean(row[1]), clean(row[2])
        board[DIMENSIONS.index(dim)].append(word)
        questions[word] = [clean(q) for q in row[3:8]]
    assert all(len(r) == 7 for r in board), "Cada dimensión debe tener 7 palabras"
    canon = {norm(w): w for r in board for w in r}

    # ── Semanas ──
    ws = wb["Frases"]
    header = [clean(c.value) for c in ws[4]]
    idx = {k: header.index(v) for k, v in FRASES_COLS.items()}
    weeks = []
    col_colors = [None] * 7
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not isinstance(row[idx["week"]], (int, float)):
            continue
        w = int(row[idx["week"]])
        get = lambda k: clean(row[idx[k]])
        phase, col = (w - 1) // 7, (3 + (w - 1) % 7) % 7
        expected = [board[(phase + k) % 7][col] for k in range(3)]
        principal = [canon[norm(get(k))] for k in ("p1", "p2", "p3")]
        assert principal == expected, f"Semana {w}: {principal} ≠ {expected}"

        emoji, _, color_name = get("color").partition(" ")
        design = get("design")
        if col_colors[col] is None:
            col_colors[col] = {"name": color_name, "emoji": emoji,
                               "palette": parse_palette(design)}
        weeks.append({
            "week": w,
            "dates": get("dates"),
            "col": col,
            "title": get("title"),
            "story": get("story"),
            "hook": get("hook"),
            "regalo": get("regalo"),
            "pensamiento": get("pensamiento"),
            "desafio": get("desafio"),
            "consejo": get("consejo"),
            "frase": get("frase"),
            "sabado": split_sabado(get("sabado")),
            "image": parse_image(design),
            "puente": puente.get(w),
        })
    assert len(weeks) == 52, f"Se esperaban 52 semanas, hay {len(weeks)}"
    assert all(col_colors), "Faltan colores de columna"

    data = {
        "year": YEAR,
        "firstSunday": FIRST_SUNDAY,
        "dimensions": DIMENSIONS,
        "board": board,
        "colors": col_colors,
        "questions": questions,
        "weeks": weeks,
    }
    OUT.write_text(
        "// Generado por tools/build_data.py desde " + SRC.name + " — no editar a mano.\n"
        "window.SUNSET_DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )
    print(f"OK: {len(weeks)} semanas, {sum(map(len, board))} palabras, "
          f"{len(puente)} semanas con Mapeo Puente → {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
